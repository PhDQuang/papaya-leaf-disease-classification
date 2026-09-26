"""Shared document object, styles and building helpers for the SV2026-08 report."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "NghienCuuKhoaHoc" / "BaoCaoTongKet_SV2026-08_FULL.docx"

doc = Document()

# ---------------------------------------------------------------- page setup
def setup_section(sec):
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    sec.left_margin, sec.right_margin = Cm(3), Cm(2)
    sec.header_distance, sec.footer_distance = Cm(1.2), Cm(1.2)


front_sec = doc.sections[0]
setup_section(front_sec)

# ------------------------------------------------------------------- styles
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(13)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.paragraph_format.line_spacing = 1.4
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.first_line_indent = Cm(0.8)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for name, size, before, after in [
    ("Title", 18, 0, 12),
    ("Heading 1", 16, 16, 10),
    ("Heading 2", 14, 12, 7),
    ("Heading 3", 13, 8, 5),
    ("Heading 4", 13, 6, 4),
]:
    s = styles[name]
    s.font.name = "Times New Roman"
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.italic = False
    s.font.color.rgb = RGBColor(0, 0, 0)
    s._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    s.paragraph_format.first_line_indent = Cm(0)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.line_spacing = 1.25
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

caption_style = styles["Caption"]
caption_style.font.name = "Times New Roman"
caption_style.font.size = Pt(12)
caption_style.font.bold = False
caption_style.font.italic = True
caption_style.font.color.rgb = RGBColor(0, 0, 0)
caption_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
caption_style.paragraph_format.first_line_indent = Cm(0)
caption_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
caption_style.paragraph_format.space_before = Pt(3)
caption_style.paragraph_format.space_after = Pt(10)
caption_style.paragraph_format.line_spacing = 1.15


# ------------------------------------------------------------- field helpers
def add_field(paragraph, instr: str, placeholder: str = ""):
    """Append a simple Word field (TOC / SEQ / PAGE) to a paragraph."""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    fld.set(qn("w:dirty"), "true")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = placeholder
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)
    return fld


def page_number_footer(sec, fmt: str = "decimal", start: int | None = None):
    sec.footer.is_linked_to_previous = False
    para = sec.footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.first_line_indent = Cm(0)
    para.paragraph_format.space_after = Pt(0)
    add_field(para, "PAGE", "1")
    for r in para.runs:
        r.font.size = Pt(12)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))
    sec._sectPr.append(pg)


page_number_footer(front_sec, "lowerRoman", 1)

# ---------------------------------------------------------------- counters
_fig_no = 0
_tab_no = 0
missing_assets: list[str] = []
# Số hiệu bảng/hình đã đăng ký, để các chương sau tham chiếu chéo mà không bị lệch số.
REF: dict[str, int] = {}


def fig_count() -> int:
    return _fig_no


def tab_count() -> int:
    return _tab_no


def ref(key: str) -> str:
    """Trả về số hiệu đã đăng ký dưới dạng chuỗi, dùng để nối vào văn bản."""
    return str(REF[key])


# ------------------------------------------------------------ text builders
def p(text: str = "", bold_lead: str | None = None, indent: bool = True, italic: bool = False):
    q = doc.add_paragraph(style="Normal")
    if not indent:
        q.paragraph_format.first_line_indent = Cm(0)
    if bold_lead and text.startswith(bold_lead):
        q.add_run(bold_lead).bold = True
        rest = q.add_run(text[len(bold_lead):])
        rest.italic = italic
    else:
        r = q.add_run(text)
        r.italic = italic
    return q


def bullet(text: str, level: int = 0):
    q = doc.add_paragraph(style="Normal")
    q.paragraph_format.first_line_indent = Cm(0)
    q.paragraph_format.left_indent = Cm(0.8 + 0.7 * level)
    q.paragraph_format.space_after = Pt(4)
    q.add_run(("– " if level == 0 else "+ ") + text)
    return q


def numbered(items: list[str]):
    for i, text in enumerate(items, 1):
        q = doc.add_paragraph(style="Normal")
        q.paragraph_format.first_line_indent = Cm(0)
        q.paragraph_format.left_indent = Cm(0.8)
        q.paragraph_format.space_after = Pt(4)
        q.add_run(f"({i}) {text}")


def formula(text: str, label: str | None = None):
    q = doc.add_paragraph(style="Normal")
    q.paragraph_format.first_line_indent = Cm(0)
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.space_before = Pt(4)
    q.paragraph_format.space_after = Pt(8)
    r = q.add_run(text)
    r.italic = True
    if label:
        q.add_run(f"     ({label})")
    return q


def code_block(lines: list[str]):
    q = doc.add_paragraph(style="Normal")
    q.paragraph_format.first_line_indent = Cm(0)
    q.paragraph_format.left_indent = Cm(0.6)
    q.paragraph_format.line_spacing = 1.0
    q.paragraph_format.space_after = Pt(8)
    q.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for i, line in enumerate(lines):
        r = q.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(10.5)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        if i < len(lines) - 1:
            r.add_break()
    return q


def h(text: str, level: int = 1):
    return doc.add_paragraph(text, style=f"Heading {level}")


def h_center(text: str, level: int = 1):
    q = doc.add_paragraph(text, style=f"Heading {level}")
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return q


def new_page():
    doc.add_page_break()


def centered(text: str, size=13, bold=False, after=8, italic=False):
    q = doc.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.first_line_indent = Cm(0)
    q.paragraph_format.space_after = Pt(after)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        r = q.add_run(line)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(size)
        if i < len(lines) - 1:
            r.add_break()
    return q


def note(text: str):
    q = doc.add_paragraph(style="Normal")
    q.paragraph_format.first_line_indent = Cm(0)
    q.paragraph_format.left_indent = Cm(0.4)
    r = q.add_run("[CẦN BỔ SUNG] " + text)
    r.italic = True
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    missing_assets.append(text)
    return q


# ---------------------------------------------------------------- captions
def _caption(prefix: str, seq_id: str, text: str, above: bool = False):
    q = doc.add_paragraph(style="Caption")
    q.paragraph_format.keep_with_next = above
    q.add_run(f"{prefix} ")
    add_field(q, f' SEQ {seq_id} \\* ARABIC ', "1")
    q.add_run(f". {text}")
    return q


# ------------------------------------------------------------------- tables
def _set_borders(t):
    for row in t.rows:
        for cell in row.cells:
            tcpr = cell._tc.get_or_add_tcPr()
            borders = tcpr.first_child_found_in("w:tcBorders")
            if borders is None:
                borders = OxmlElement("w:tcBorders")
                tcpr.append(borders)
            for side in ("top", "left", "bottom", "right"):
                el = OxmlElement(f"w:{side}")
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "4")
                el.set(qn("w:color"), "808080")
                borders.append(el)


def table(caption: str, header: list[str], rows: list[list[str]],
          widths: list[float] | None = None, font_size: float = 11,
          key: str | None = None) -> int:
    """Insert a captioned table. Returns its sequential number."""
    global _tab_no
    _tab_no += 1
    if key:
        REF[key] = _tab_no
    cap = _caption("Bảng", "Bang", caption, above=True)
    cap.paragraph_format.space_before = Pt(8)
    cap.paragraph_format.space_after = Pt(4)
    for r in cap.runs:
        r.italic = False
        r.bold = True

    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, value in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = value
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            c.width = Cm(widths[i])
        for pp in c.paragraphs:
            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pp.paragraph_format.first_line_indent = Cm(0)
            pp.paragraph_format.space_after = Pt(2)
            pp.paragraph_format.line_spacing = 1.1
            for rr in pp.runs:
                rr.bold = True
                rr.font.size = Pt(font_size)
        tcpr = c._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "E8E8E8")
        tcpr.append(shd)
    trpr = t.rows[0]._tr.get_or_add_trPr()
    header_el = OxmlElement("w:tblHeader")
    header_el.set(qn("w:val"), "true")
    trpr.append(header_el)

    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cells[i].width = Cm(widths[i])
            for pp in cells[i].paragraphs:
                pp.paragraph_format.first_line_indent = Cm(0)
                pp.paragraph_format.space_after = Pt(1)
                pp.paragraph_format.line_spacing = 1.1
                if i > 0:
                    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for rr in pp.runs:
                    rr.font.size = Pt(font_size)
    _set_borders(t)
    tail = doc.add_paragraph()
    tail.paragraph_format.space_after = Pt(2)
    tail.paragraph_format.line_spacing = 1.0
    return _tab_no


# ------------------------------------------------------------------ figures
def figure(rel: str, caption: str, width: float = 14.5) -> int:
    """Insert a captioned picture. Always consumes a figure number so that
    cross-references written in the prose stay correct even if a file is
    still missing."""
    global _fig_no
    _fig_no += 1
    path = ROOT / rel
    if path.exists():
        q = doc.add_paragraph()
        q.paragraph_format.first_line_indent = Cm(0)
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.keep_with_next = True
        q.paragraph_format.space_before = Pt(6)
        q.paragraph_format.space_after = Pt(2)
        q.add_run().add_picture(str(path), width=Cm(width))
    else:
        note(f"Thiếu ảnh cho Hình {_fig_no} ({caption}). Đường dẫn dự kiến: {rel}. "
             f"Hãy chèn ảnh vào đúng vị trí này.")
    _caption("Hình", "Hinh", caption)
    return _fig_no


def figure_row(items: list[tuple[str, str]], caption: str, width: float = 4.4) -> int:
    """Insert several small pictures on one line, then a single caption."""
    global _fig_no
    _fig_no += 1
    q = doc.add_paragraph()
    q.paragraph_format.first_line_indent = Cm(0)
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.keep_with_next = True
    q.paragraph_format.space_before = Pt(6)
    q.paragraph_format.space_after = Pt(2)
    missing = []
    for rel, label in items:
        path = ROOT / rel
        if path.exists():
            run = q.add_run()
            run.add_picture(str(path), width=Cm(width))
            q.add_run("  ")
        else:
            missing.append(rel)
    if missing:
        note(f"Thiếu ảnh trong Hình {_fig_no}: {', '.join(missing)}.")
    sub = doc.add_paragraph()
    sub.paragraph_format.first_line_indent = Cm(0)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(0)
    sub.paragraph_format.keep_with_next = True
    r = sub.add_run("   ".join(f"({chr(97 + i)}) {lbl}" for i, (_, lbl) in enumerate(items)))
    r.font.size = Pt(11)
    _caption("Hình", "Hinh", caption)
    return _fig_no


# ------------------------------------------------------------------ closing
def start_body_section():
    """Open the numbered body section (arabic page numbers restarting at 1)."""
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    setup_section(sec)
    page_number_footer(sec, "decimal", 1)
    return sec


def save(path: Path | None = None) -> Path:
    """Lưu tài liệu, bật cờ để Word tự cập nhật các trường TOC/SEQ/PAGE khi mở.

    Truyền `path` khi muốn xuất một trích đoạn ra tệp riêng thay vì ghi đè báo cáo chính.
    """
    target = Path(path) if path is not None else OUT
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.append(upd)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    return target
