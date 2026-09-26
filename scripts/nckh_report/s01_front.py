"""Bìa chính, bìa phụ, mục lục, danh mục bảng/hình/chữ viết tắt, PL03."""
from __future__ import annotations

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from .core import (add_field, bullet, centered, doc, h_center, new_page, note, p,
                   start_body_section)

TITLE = "PHÁT HIỆN VÀ PHÂN LOẠI BỆNH Ở LÁ CÂY ĐU ĐỦ\nSỬ DỤNG MÔ HÌNH YOLOv11 VÀ CNN"


def _cover(inner: bool) -> None:
    centered("BỘ GIÁO DỤC VÀ ĐÀO TẠO", 14, True, 2)
    centered("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ KỸ THUẬT", 14, True, 2)
    centered("THÀNH PHỐ HỒ CHÍ MINH", 14, True, 30)
    centered("BÁO CÁO TỔNG KẾT", 18, True, 6)
    centered("ĐỀ TÀI NGHIÊN CỨU KHOA HỌC CỦA SINH VIÊN", 15, True, 40)
    centered(TITLE, 17, True, 10)
    centered("Mã số đề tài: SV2026-08", 14, True, 34)
    if inner:
        centered("Thuộc nhóm ngành khoa học: Công nghệ thông tin", 13, False, 18)
        block = [
            "Sinh viên chịu trách nhiệm chính: Phạm Đăng Quang       Nam, Nữ: Nam",
            "Mã số sinh viên: 23110143",
            "Dân tộc: Kinh",
            "Lớp, khoa: ................ , Khoa Công nghệ thông tin       Năm thứ: 3 / Số năm đào tạo: 4",
            "Ngành học: Công nghệ thông tin",
            "Thành viên: Lê Quang Sang (MSSV 23110701); Vũ Minh Hiếu (MSSV 23110218)",
            "Người hướng dẫn: TS. Bùi Mạnh Quân",
        ]
        for line in block:
            q = p(line, indent=False)
            q.alignment = WD_ALIGN_PARAGRAPH.LEFT
            q.paragraph_format.space_after = Pt(3)
        note("Điền chính xác lớp, dân tộc, năm thứ và nhóm ngành khoa học theo hồ sơ "
             "đã được Phòng KHCN phê duyệt trước khi in nộp.")
    else:
        centered("Chủ nhiệm đề tài: Phạm Đăng Quang", 14, False, 8)
        centered("Người hướng dẫn: TS. Bùi Mạnh Quân", 14, False, 30)
    centered("TP. Hồ Chí Minh, tháng 9 năm 2026", 13, False, 0)


def _toc_block(title: str, instr: str, hint: str) -> None:
    h_center(title, 1)
    q = p("", indent=False)
    q.paragraph_format.space_after = Pt(0)
    add_field(q, instr, hint)


ABBREV = [
    ("AI", "Artificial Intelligence", "Trí tuệ nhân tạo"),
    ("AMP", "Automatic Mixed Precision", "Huấn luyện độ chính xác hỗn hợp tự động"),
    ("AP", "Average Precision", "Độ chính xác trung bình"),
    ("API", "Application Programming Interface", "Giao diện lập trình ứng dụng"),
    ("BCE", "Binary Cross-Entropy", "Hàm mất mát entropy chéo nhị phân"),
    ("CBAM", "Convolutional Block Attention Module", "Khối chú ý tích chập"),
    ("CNN", "Convolutional Neural Network", "Mạng nơ-ron tích chập"),
    ("CIoU", "Complete Intersection over Union", "Hàm mất mát IoU đầy đủ"),
    ("CPU", "Central Processing Unit", "Bộ xử lý trung tâm"),
    ("DFL", "Distribution Focal Loss", "Hàm mất mát tiêu điểm phân phối"),
    ("FN", "False Negative", "Âm tính giả"),
    ("FP", "False Positive", "Dương tính giả"),
    ("FPS", "Frames Per Second", "Số khung hình xử lý mỗi giây"),
    ("GPU", "Graphics Processing Unit", "Bộ xử lý đồ họa"),
    ("IoU", "Intersection over Union", "Tỷ lệ giao trên hợp"),
    ("mAP", "mean Average Precision", "Độ chính xác trung bình toàn cục"),
    ("MBConv", "Mobile Inverted Bottleneck Convolution", "Khối tích chập nút cổ chai đảo"),
    ("NMS", "Non-Maximum Suppression", "Loại bỏ khung trùng lặp"),
    ("ReLU", "Rectified Linear Unit", "Hàm kích hoạt tuyến tính chỉnh lưu"),
    ("REST", "Representational State Transfer", "Kiến trúc dịch vụ web REST"),
    ("ROI", "Region of Interest", "Vùng quan tâm"),
    ("SE", "Squeeze-and-Excitation", "Khối nén và kích thích"),
    ("SPPF", "Spatial Pyramid Pooling – Fast", "Gộp kim tự tháp không gian phiên bản nhanh"),
    ("SVM", "Support Vector Machine", "Máy vec-tơ hỗ trợ"),
    ("TN", "True Negative", "Âm tính thật"),
    ("TP", "True Positive", "Dương tính thật"),
    ("ViT", "Vision Transformer", "Mô hình Transformer cho thị giác máy tính"),
    ("YOLO", "You Only Look Once", "Họ mô hình phát hiện đối tượng một giai đoạn"),
]


def build() -> None:
    _cover(False)
    new_page()
    _cover(True)
    new_page()

    _toc_block("MỤC LỤC", ' TOC \\o "1-3" \\h \\z \\u ',
               "Nhấn chuột phải vào vùng này và chọn Update Field để sinh mục lục.")
    new_page()

    _toc_block("DANH MỤC BẢNG BIỂU", ' TOC \\h \\z \\c "Bang" ',
               "Nhấn chuột phải vào vùng này và chọn Update Field để sinh danh mục bảng.")
    new_page()

    _toc_block("DANH MỤC HÌNH ẢNH", ' TOC \\h \\z \\c "Hinh" ',
               "Nhấn chuột phải vào vùng này và chọn Update Field để sinh danh mục hình.")
    new_page()

    h_center("DANH MỤC NHỮNG TỪ VIẾT TẮT", 1)
    from .core import table  # local import to keep the numbering order natural
    table("Danh mục các từ viết tắt sử dụng trong báo cáo",
          ["Viết tắt", "Thuật ngữ tiếng Anh", "Nghĩa tiếng Việt"],
          [[a, b, c] for a, b, c in ABBREV],
          widths=[2.4, 6.6, 7.0], font_size=11)
    new_page()

    _research_info()


def _research_info() -> None:
    centered("BỘ GIÁO DỤC VÀ ĐÀO TẠO", 12, True, 0)
    centered("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ KỸ THUẬT TP. HỒ CHÍ MINH", 12, True, 14)
    h_center("THÔNG TIN KẾT QUẢ NGHIÊN CỨU CỦA ĐỀ TÀI", 1)

    p("1. Thông tin chung:", bold_lead="1. Thông tin chung:", indent=False)
    for line in [
        "Tên đề tài: Phát hiện và phân loại bệnh ở lá cây đu đủ sử dụng mô hình YOLOv11 và CNN.",
        "Mã số đề tài: SV2026-08.",
        "Chủ nhiệm đề tài: Phạm Đăng Quang — MSSV 23110143 — SĐT 0347560113 — "
        "Email 23110143@student.hcmute.edu.vn.",
        "Thành viên: Lê Quang Sang (MSSV 23110701); Vũ Minh Hiếu (MSSV 23110218).",
        "Người hướng dẫn: TS. Bùi Mạnh Quân — Khoa Công nghệ thông tin — Email quanbm@hcmute.edu.vn.",
        "Thời gian thực hiện: 12 tháng.",
        "Kinh phí được duyệt: 10.000.000 đồng.",
    ]:
        bullet(line)

    p("2. Mục tiêu đề tài:", bold_lead="2. Mục tiêu đề tài:", indent=False)
    p("Xây dựng hệ thống tự động phát hiện và phân loại các loại bệnh hại trên lá cây đu đủ "
      "dựa trên công nghệ trí tuệ nhân tạo với mô hình YOLOv11 và mạng nơ-ron tích chập. "
      "Hệ thống nhằm hỗ trợ nông dân nhận biết sớm dịch bệnh, từ đó đưa ra biện pháp xử lý "
      "kịp thời, giúp giảm thiểu thiệt hại kinh tế và nâng cao năng suất cây trồng.")

    p("3. Tính mới và sáng tạo:", bold_lead="3. Tính mới và sáng tạo:", indent=False)
    for line in [
        "Đề xuất kiến trúc xếp tầng hai giai đoạn YOLOv11m → EfficientNet-B2, trong đó giai đoạn "
        "một khoanh vùng tổn thương và giai đoạn hai phân loại chi tiết trên từng vùng quan tâm, "
        "thay vì phân loại trực tiếp trên toàn ảnh như đa số công trình đã công bố.",
        "Xây dựng cơ chế tổng hợp phiếu bầu có trọng số ở mức ảnh, kết hợp độ tin cậy của bộ phát "
        "hiện và xác suất của bộ phân loại, giúp giảm ảnh hưởng của các khung nhiễu.",
        "Thiết kế cơ chế hạ ngưỡng tin cậy thích nghi kèm nhánh dự phòng phân loại toàn ảnh, bảo "
        "đảm hệ thống luôn trả về kết quả ngay cả khi bộ phát hiện không tìm thấy tổn thương.",
        "Đóng gói toàn bộ mô hình thành dịch vụ REST bằng FastAPI, giao diện web bằng Gradio và "
        "container Docker, cho phép tái lập kết quả và triển khai nhanh.",
    ]:
        bullet(line)

    p("4. Kết quả nghiên cứu:", bold_lead="4. Kết quả nghiên cứu:", indent=False)
    for line in [
        "Bộ phát hiện YOLOv11m đạt mAP@0,5 = 0,7767 và mAP@0,5:0,95 = 0,6520 trên tập kiểm tra.",
        "Bộ phân loại EfficientNet-B2 trên vùng quan tâm đạt độ chính xác 96,02% với 1.836 mẫu kiểm tra.",
        "Hệ thống xếp tầng hoàn chỉnh đạt độ chính xác 94,85% và F1 trung bình vĩ mô 94,92% trên "
        "291 ảnh kiểm tra ở mức ảnh, cao hơn mô hình phân loại toàn ảnh cùng kiến trúc 2,07 điểm phần trăm.",
        "Ba cấu hình tham chiếu ResNet50, DenseNet121 và VGG16 ghép cùng YOLOv11m đều cho kết quả "
        "thấp hơn cấu hình đề xuất, xác nhận lựa chọn EfficientNet-B2 là hợp lý.",
    ]:
        bullet(line)

    p("5. Đóng góp về mặt kinh tế – xã hội, giáo dục và đào tạo, an ninh, quốc phòng "
      "và khả năng áp dụng của đề tài:",
      bold_lead="5. Đóng góp về mặt kinh tế – xã hội, giáo dục và đào tạo, an ninh, quốc phòng "
                "và khả năng áp dụng của đề tài:", indent=False)
    p("Kết quả của đề tài cung cấp một công cụ chẩn đoán bệnh lá đu đủ nhanh, chi phí thấp, chỉ "
      "yêu cầu một bức ảnh chụp bằng điện thoại. Công cụ này hỗ trợ nông hộ phát hiện sớm các "
      "bệnh thán thư, đốm vòng, đốm vi khuẩn và xoăn lá, qua đó giảm thiệt hại năng suất và chi "
      "phí thuốc bảo vệ thực vật sử dụng sai mục đích. Về đào tạo, mã nguồn, sổ tay thực nghiệm "
      "và bộ kết quả được công bố kèm theo có thể dùng làm học liệu cho các học phần thị giác "
      "máy tính và học sâu tại Trường.")

    p("6. Công bố khoa học của sinh viên từ kết quả nghiên cứu của đề tài:",
      bold_lead="6. Công bố khoa học của sinh viên từ kết quả nghiên cứu của đề tài:", indent=False)
    note("Điền tên đầy đủ hội nghị/tạp chí, mã số ISBN/ISSN, ngày nhận đăng, đường dẫn bài báo và "
         "đính kèm giấy chấp nhận đăng (bài báo mã số GTSD2026-148, đã được chấp nhận đăng, có "
         "lời cảm ơn ghi rõ Grant No. SV2026-08). Hồ sơ minh chứng đặt tại Phụ lục A.")

    p("")
    q = p("Ngày ......... tháng ......... năm 2026", indent=False)
    q.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    tbl = doc.add_table(rows=1, cols=2)
    tbl.autofit = False
    left, right = tbl.rows[0].cells
    left.width, right.width = Cm(8), Cm(8)
    for cell, lines in (
        (left, ["Người hướng dẫn", "(ký, họ và tên)", "", "", "", "TS. Bùi Mạnh Quân"]),
        (right, ["Chủ nhiệm đề tài", "(ký, họ và tên)", "", "", "", "Phạm Đăng Quang"]),
    ):
        cell.text = ""
        for i, line in enumerate(lines):
            para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.first_line_indent = Cm(0)
            para.paragraph_format.space_after = Pt(0)
            r = para.add_run(line)
            r.bold = i == 0
            r.italic = i == 1
            r.font.size = Pt(13)

    start_body_section()


build()
